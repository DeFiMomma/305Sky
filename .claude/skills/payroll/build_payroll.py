#!/usr/bin/env python3
"""Build the 305 SKY payroll workbook for one pay period.

Usage:
    python3 .claude/skills/payroll/build_payroll.py payroll/<period-start>/hours.csv \
        [--rates .claude/skills/payroll/rates.csv] [--out payroll/<period-start>/payroll.xlsx]

hours.csv columns: employee, location, one column per date (YYYY-MM-DD, 7 or 14 days,
starting on the Friday the workweek starts), and an optional report_worked column holding
the "Worked" total from the time-clock report (used only as a cross-check).

Overtime is figured per 7-day workweek: hours over 40 in a week are paid at 1.5x. Rows for the
same employee at different locations are combined before overtime is figured.
"""
import argparse
import csv
import sys
from collections import OrderedDict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OT_THRESHOLD = 40.0
OT_MULTIPLIER = 1.5

FONT = "Arial"
BLUE = "0000FF"
GREEN = "008000"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEADER_FILL = PatternFill("solid", fgColor="1F1F1F")
TOTAL_FILL = PatternFill("solid", fgColor="E7E6E6")
THIN = Side(style="thin", color="BFBFBF")
MONEY = '$#,##0.00;($#,##0.00);"-"'
HOURS = '0.00;(0.00);"-"'


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def read_rates(path):
    rates = OrderedDict()
    with open(path, newline="") as fh:
        for row in csv.DictReader(fh):
            name = row["employee"].strip()
            rates[name] = {
                "rate": float(row["rate"]) if row["rate"].strip() else None,
                "withhold": float(row.get("withhold_hours") or 0),
                "note": (row.get("note") or "").strip(),
            }
    return rates


def read_hours(path):
    with open(path, newline="") as fh:
        reader = csv.DictReader(fh)
        dates = [c for c in reader.fieldnames if c[:4].isdigit()]
        rows = []
        for row in reader:
            rows.append({
                "employee": row["employee"].strip(),
                "location": (row.get("location") or "").strip(),
                "hours": [float(row[d]) if row[d].strip() else 0.0 for d in dates],
                "report_worked": float(row["report_worked"]) if (row.get("report_worked") or "").strip() else None,
            })
    dates = [date.fromisoformat(d) for d in dates]
    if len(dates) not in (7, 14):
        sys.exit(f"Expected 7 or 14 date columns, found {len(dates)}")
    return dates, rows


def compute(dates, rows, rates):
    """Plain-Python payroll, used to cross-check the workbook formulas."""
    weeks = len(dates) // 7
    people = OrderedDict()
    for r in rows:
        p = people.setdefault(r["employee"], [0.0] * weeks)
        for w in range(weeks):
            p[w] += sum(r["hours"][w * 7:(w + 1) * 7])
    result = OrderedDict()
    for name, week_hours in people.items():
        info = rates.get(name, {"rate": None, "withhold": 0.0})
        lines = []
        for w, hrs in enumerate(week_hours):
            withheld = info["withhold"] if w == 0 else 0.0
            payable = max(hrs - withheld, 0.0)
            reg = min(payable, OT_THRESHOLD)
            ot = max(payable - OT_THRESHOLD, 0.0)
            gross = None if info["rate"] is None else round(reg * info["rate"] + ot * info["rate"] * OT_MULTIPLIER, 2)
            lines.append(dict(week=w + 1, worked=hrs, withheld=withheld, reg=reg, ot=ot, gross=gross))
        result[name] = lines
    return result


def build(dates, rows, rates, out):
    weeks = len(dates) // 7
    start, end = dates[0], dates[-1]
    period = f"{start:%b} {start.day} – {end:%b} {end.day}, {end.year}"
    wb = Workbook()

    # ---------------- Rates ----------------
    ws_r = wb.active
    ws_r.title = "Rates"
    ws_r.append(["Employee", "Hourly Rate", "Hours Withheld (this period)", "Note"])
    employees = list(OrderedDict.fromkeys(r["employee"] for r in rows))
    rate_names = list(OrderedDict.fromkeys(list(rates) + employees))
    for name in rate_names:
        info = rates.get(name, {"rate": None, "withhold": 0.0, "note": "RATE NOT ON RATE SHEET"})
        ws_r.append([name, info["rate"], info["withhold"], info["note"]])
    for row in ws_r.iter_rows(min_row=2, max_row=ws_r.max_row):
        row[0].font = f()
        for c in (row[1], row[2]):
            c.font = f(color=BLUE)
        row[1].number_format = MONEY
        row[2].number_format = HOURS
        row[3].font = f(italic=True)
        if row[1].value is None:
            row[1].fill = YELLOW
        if row[2].value:
            row[2].fill = YELLOW
    rates_last = ws_r.max_row
    style_header(ws_r, 1, 4)
    widths(ws_r, [24, 14, 26, 60])

    # ---------------- Hours Detail ----------------
    ws_h = wb.create_sheet("Hours Detail")
    head = ["Employee", "Location"] + [f"{d:%a} {d.month}/{d.day}" for d in dates]
    head += [f"Week {w + 1} Hrs" for w in range(weeks)] + ["Report Worked", "Check"]
    ws_h.append(head)
    for i, r in enumerate(rows, start=2):
        ws_h.cell(i, 1, r["employee"]).font = f()
        ws_h.cell(i, 2, r["location"] or None).font = f()
        for j, h in enumerate(r["hours"]):
            c = ws_h.cell(i, 3 + j, h if h else None)
            c.font = f(color=BLUE)
            c.number_format = HOURS
        for w in range(weeks):
            a = get_column_letter(3 + w * 7)
            b = get_column_letter(3 + w * 7 + 6)
            c = ws_h.cell(i, 3 + len(dates) + w, f"=SUM({a}{i}:{b}{i})")
            c.font = f(bold=True)
            c.number_format = HOURS
        rep_col = 3 + len(dates) + weeks
        c = ws_h.cell(i, rep_col, r["report_worked"])
        c.font = f(color=BLUE)
        c.number_format = HOURS
        first_wk = get_column_letter(3 + len(dates))
        last_wk = get_column_letter(2 + len(dates) + weeks)
        rep = get_column_letter(rep_col)
        c = ws_h.cell(i, rep_col + 1,
                      f'=IF({rep}{i}="","",IF(ABS(SUM({first_wk}{i}:{last_wk}{i})-{rep}{i})<0.05,"OK","MISMATCH"))')
        c.font = f()
    hours_last = ws_h.max_row
    style_header(ws_h, 1, len(head))
    widths(ws_h, [24, 10] + [9] * len(dates) + [12] * weeks + [14, 11])
    ws_h.freeze_panes = "C2"

    # ---------------- Payroll ----------------
    ws = wb.create_sheet("Payroll", 0)
    ws["A1"] = "305 SKY LLC — Payroll"
    ws["A1"].font = f(bold=True, size=14)
    ws["A2"] = f"Pay period: {period}  ({'Weekly' if weeks == 1 else 'Bi-weekly'}; workweek runs Fri–Thu)"
    ws["A2"].font = f(italic=True)
    ws["A4"] = "Overtime after (hrs / week)"
    ws["B4"] = OT_THRESHOLD
    ws["A5"] = "Overtime multiplier"
    ws["B5"] = OT_MULTIPLIER
    for c in ("A4", "A5"):
        ws[c].font = f()
    for c in ("B4", "B5"):
        ws[c].font = f(color=BLUE)
        ws[c].fill = YELLOW
    ws["B4"].number_format = "0.0"
    ws["B5"].number_format = '0.0"x"'

    cols = ["Employee", "Week", "Rate", "Hours Worked", "Hours Withheld", "Payable Hrs",
            "Regular Hrs", "OT Hrs", "Regular Pay", "OT Pay", "Gross Pay", "Flag"]
    hr = 7
    for j, name in enumerate(cols, start=1):
        ws.cell(hr, j, name)
    style_header(ws, hr, len(cols))

    rng = lambda col: f"'Hours Detail'!${col}$2:${col}${hours_last}"
    rr = lambda col: f"Rates!${col}$2:${col}${rates_last}"
    r = hr + 1
    for name in employees:
        for w in range(weeks):
            wk_col = get_column_letter(3 + len(dates) + w)
            ws.cell(r, 1, name).font = f()
            ws.cell(r, 2, w + 1).font = f()
            ws.cell(r, 3, f"=IFERROR(IF(INDEX({rr('B')},MATCH(A{r},{rr('A')},0))=\"\",\"\","
                          f"INDEX({rr('B')},MATCH(A{r},{rr('A')},0))),\"\")").font = f(color=GREEN)
            ws.cell(r, 4, f"=ROUND(SUMIFS({rng(wk_col)},{rng('A')},A{r}),2)").font = f(color=GREEN)
            if w == 0:
                ws.cell(r, 5, f"=IFERROR(INDEX({rr('C')},MATCH(A{r},{rr('A')},0)),0)").font = f(color=GREEN)
            else:
                ws.cell(r, 5, 0).font = f(color=BLUE)
            ws.cell(r, 6, f"=ROUND(MAX(D{r}-E{r},0),2)").font = f()
            ws.cell(r, 7, f"=MIN(F{r},$B$4)").font = f()
            ws.cell(r, 8, f"=ROUND(MAX(F{r}-$B$4,0),2)").font = f()
            ws.cell(r, 9, f'=IF(C{r}="","",ROUND(G{r}*C{r},2))').font = f()
            ws.cell(r, 10, f'=IF(C{r}="","",ROUND(H{r}*C{r}*$B$5,2))').font = f()
            ws.cell(r, 11, f'=IF(C{r}="","",I{r}+J{r})').font = f(bold=True)
            ws.cell(r, 12, f'=IF(C{r}="","RATE NEEDED",IF(E{r}>0,"Hours withheld — confirm",""))').font = f(bold=True, color="C00000")
            for j in (4, 5, 6, 7, 8):
                ws.cell(r, j).number_format = HOURS
            for j in (3, 9, 10, 11):
                ws.cell(r, j).number_format = MONEY
            for j in range(1, len(cols) + 1):
                ws.cell(r, j).border = Border(bottom=THIN)
            r += 1
    last = r - 1
    ws.cell(r, 1, "TOTAL").font = f(bold=True)
    for j in (4, 5, 6, 7, 8, 9, 10, 11):
        L = get_column_letter(j)
        c = ws.cell(r, j, f"=SUM({L}{hr + 1}:{L}{last})")
        c.font = f(bold=True)
        c.number_format = MONEY if j >= 9 else HOURS
    ws.cell(r, 12, f'=IF(COUNTIF(L{hr + 1}:L{last},"RATE NEEDED")>0,'
                   f'COUNTIF(L{hr + 1}:L{last},"RATE NEEDED")&" rate(s) missing — total incomplete","")').font = f(bold=True, color="C00000")
    for j in range(1, len(cols) + 1):
        ws.cell(r, j).fill = TOTAL_FILL
    total_row = r

    notes = [
        "How it works:",
        "• Hours come from the 'Hours Detail' tab (time-clock Payroll Report). Rows for the same person at different locations are added together before overtime.",
        "• Overtime: hours over B4 in each Fri–Thu workweek are paid at the B5 multiplier. Overtime is never averaged across the two weeks of a bi-weekly period.",
        "• Rates and withheld hours come from the 'Rates' tab (from the 2026 rate sheet). Yellow cells need your input; blue text = typed-in values.",
        "• Gross pay is before taxes and deductions — QuickBooks figures those when you run payroll.",
    ]
    for k, line in enumerate(notes):
        c = ws.cell(total_row + 2 + k, 1, line)
        c.font = f(bold=(k == 0), italic=(k > 0))
    widths(ws, [24, 7, 11, 13, 14, 12, 12, 9, 13, 12, 13, 30])
    ws.freeze_panes = ws.cell(hr + 1, 2)

    # ---------------- QuickBooks entry ----------------
    ws_q = wb.create_sheet("QuickBooks Entry", 1)
    ws_q["A1"] = f"Enter in QuickBooks Payroll — {period}"
    ws_q["A1"].font = f(bold=True, size=14)
    ws_q.append([])
    ws_q.append(["Employee", "Regular Hours", "Overtime Hours", "Rate", "Gross Pay"])
    style_header(ws_q, 3, 5)
    pr = lambda col: f"Payroll!${col}${hr + 1}:${col}${last}"
    for i, name in enumerate(employees, start=4):
        ws_q.cell(i, 1, name).font = f()
        ws_q.cell(i, 2, f"=SUMIFS({pr('G')},{pr('A')},A{i})").font = f(color=GREEN)
        ws_q.cell(i, 3, f"=SUMIFS({pr('H')},{pr('A')},A{i})").font = f(color=GREEN)
        ws_q.cell(i, 4, f"=INDEX({pr('C')},MATCH(A{i},{pr('A')},0))").font = f(color=GREEN)
        ws_q.cell(i, 5, f'=IF(D{i}="","RATE NEEDED",SUMIFS({pr("K")},{pr("A")},A{i}))').font = f(bold=True, color=GREEN)
        ws_q.cell(i, 2).number_format = HOURS
        ws_q.cell(i, 3).number_format = HOURS
        ws_q.cell(i, 4).number_format = MONEY
        ws_q.cell(i, 5).number_format = MONEY
    q_last = 3 + len(employees)
    t = q_last + 1
    ws_q.cell(t, 1, "TOTAL").font = f(bold=True)
    for j in (2, 3, 5):
        L = get_column_letter(j)
        c = ws_q.cell(t, j, f"=SUM({L}4:{L}{q_last})")
        c.font = f(bold=True)
        c.number_format = MONEY if j == 5 else HOURS
        ws_q.cell(t, j).fill = TOTAL_FILL
    ws_q.cell(t, 1).fill = TOTAL_FILL
    ws_q.cell(t, 4).fill = TOTAL_FILL
    ws_q.cell(t + 2, 1, "QuickBooks multiplies hours by the rate saved on each employee's profile — make sure those "
                        "rates match column D.").font = f(italic=True)
    widths(ws_q, [24, 15, 15, 11, 14])

    wb.save(out)


def style_header(ws, row, ncols):
    for j in range(1, ncols + 1):
        c = ws.cell(row, j)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, ws_widths):
    for j, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


def main():
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser()
    ap.add_argument("hours")
    ap.add_argument("--rates", default=str(here / "rates.csv"))
    ap.add_argument("--out")
    a = ap.parse_args()

    rates = read_rates(a.rates)
    dates, rows = read_hours(a.hours)
    if dates[0].weekday() != 4:
        print(f"WARNING: period starts {dates[0]:%A}, but the workweek starts on Friday.")
    out = a.out or str(Path(a.hours).with_name(f"305SKY_Payroll_{dates[0]}_to_{dates[-1]}.xlsx"))

    # Cross-checks against the time-clock report
    for r in rows:
        if r["report_worked"] is not None and abs(sum(r["hours"]) - r["report_worked"]) > 0.05:
            print(f"WARNING: {r['employee']} {r['location']}: daily hours sum to {sum(r['hours']):.1f}, "
                  f"report says {r['report_worked']:.1f}")
    unknown = sorted({r["employee"] for r in rows} - set(rates))
    for name in unknown:
        print(f"WARNING: {name} is not in rates.csv")

    build(dates, rows, rates, out)

    result = compute(dates, rows, rates)
    print(f"\nWrote {out}\n")
    print(f"{'Employee':<20}{'Wk':>3}{'Worked':>8}{'Withh':>7}{'Reg':>7}{'OT':>6}{'Rate':>8}{'Gross':>11}")
    total = 0.0
    missing = []
    for name, lines in result.items():
        rate = rates.get(name, {}).get("rate")
        for l in lines:
            g = "RATE NEEDED" if l["gross"] is None else f"${l['gross']:,.2f}"
            print(f"{name:<20}{l['week']:>3}{l['worked']:>8.1f}{l['withheld']:>7.1f}{l['reg']:>7.1f}{l['ot']:>6.1f}"
                  f"{('$%.2f' % rate) if rate else '—':>8}{g:>11}")
            if l["gross"] is None:
                missing.append(name)
            else:
                total += l["gross"]
    print(f"\nGross payroll (employees with rates): ${total:,.2f}")
    if missing:
        print(f"Missing rates: {', '.join(sorted(set(missing)))}")


if __name__ == "__main__":
    main()
